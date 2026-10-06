#!/usr/bin/env python3
"""Word version of the manuscript, built from the same sources as the PDF (main.tex, the generated
tables, the figures, and the label, citation and bibliography data of the last LaTeX run).

    latexmk -pdf main.tex          # writes main.aux and main.bbl, which this script reads
    python3 make_docx.py           # writes main.docx

Section, table, figure, equation and reference numbers are taken from main.aux, so they are the
PDF's. Equations become native Word equations (pandoc's LaTeX reader). Requires pandoc and
python-docx."""
import pathlib
import re
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = pathlib.Path(__file__).resolve().parent
TEX, AUX, BBL, OUT = HERE / "main.tex", HERE / "main.aux", HERE / "main.bbl", HERE / "main.docx"
FONT = "Times New Roman"


def braced(s, i):
    """Return (content, end) of the brace group starting at s[i] == '{'."""
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while True:
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1


def labels_and_cites():
    aux = AUX.read_text()
    labels = {}
    for m in re.finditer(r"\\newlabel\{([^}]*)\}\{\{([^}]*)\}", aux):
        labels[m.group(1)] = m.group(2)
    # natbib author-year entries: \bibcite{key}{{n}{year}{{short author}}{{full author list}}}
    cites = {}
    for m in re.finditer(r"\\bibcite\{([^}]*)\}\{\{(\d+)\}\{([^}]*)\}\{\{([^}]*)\}\}", aux):
        cites[m.group(1)] = (int(m.group(2)), m.group(4).replace("~", " "), m.group(3))
    return labels, cites


def cite_text(keys, cites, textual=False):
    """Author-year citations as the elsarticle-harv style prints them: (Huh et al., 2025; Page, 1954)
    in parentheses, or Huh et al. (2025) when the authors are part of the sentence."""
    items = [cites[k.strip()] for k in keys.split(",")]
    if textual:
        return "; ".join(f"{a} ({y})" for _, a, y in items)
    return "(" + "; ".join(f"{a}, {y}" for _, a, y in items) + ")"


def table_rows(path):
    rows = (HERE / path).read_text()
    # Table 8 stacks each interval under its mean (a nested tabular); in Word it follows the mean
    rows = re.sub(r"\\begin\{tabular\}\[t\]\{@\{\}r@\{\}\}(.*?)\\\\\{\\scriptsize (\[.*?\])\}\\end\{tabular\}", r"\1 \2", rows)
    return rows.replace("\\\\[1pt]", "\\\\")


def simplify_colspec(spec):
    spec = re.sub(r">\{[^{}]*(\{[^{}]*\}[^{}]*)*\}", "", spec)       # >{\raggedright\arraybackslash}
    spec = re.sub(r"@\{[^{}]*\}", "", spec)
    spec = re.sub(r"p\{[^{}]*\}", "l", spec)
    return spec


def body_latex():
    s = TEX.read_text()
    labels, cites = labels_and_cites()
    s = s[s.index("\\begin{document}") + len("\\begin{document}"):s.index("\\end{document}")]

    # front matter -> title, author lines, abstract, keywords
    fm = s[s.index("\\begin{frontmatter}"):s.index("\\end{frontmatter}") + len("\\end{frontmatter}")]
    title, _ = braced(fm, fm.index("\\title{") + 6)
    org = re.search(r"organization=\{([^}]*)\}", fm).group(1)
    city = re.search(r"city=\{([^}]*)\}", fm)
    country = re.search(r"country=\{([^}]*)\}", fm).group(1)
    if city:
        country = f"{city.group(1)}, {country}"
    email = re.search(r"\\ead\{([^}]*)\}", fm).group(1)
    author = re.search(r"\\author\[[^\]]*\]\{([^\\}]*)", fm).group(1).strip()
    abstract = fm[fm.index("\\begin{abstract}") + 16:fm.index("\\end{abstract}")].strip()
    keywords = re.sub(r"\s*\\sep\s*", "; ", fm[fm.index("\\begin{keyword}") + 15:fm.index("\\end{keyword}")].strip())
    front = (f"TITLEMARK {title}\n\nAUTHORMARK {author}\\textsuperscript{{*}}\n\n"
             f"AFFILMARK {org}, {country}\n\nCORRMARK \\textsuperscript{{*}}Corresponding author. E-mail: {email}\n\n"
             f"ABSTRACTHEAD Abstract\n\nABSTRACTMARK {abstract}\n\nKEYWORDSMARK \\textit{{Keywords:}} {keywords}\n\n")
    s = s.replace(fm, front)

    # generated tables, figures, layout-only commands
    s = re.sub(r"\\tablerows\{([^}]*)\}", lambda m: table_rows(m.group(1)), s)
    # tabular* filled to the column or text width -> a plain tabular (Word tables are fitted below)
    s = re.sub(r"\\begin\{tabular\*\}\{\\(?:textwidth|columnwidth)\}\{@\{\\extracolsep\{\\fill\}\}", lambda m: "\\begin{tabular}{", s)
    s = s.replace("\\end{tabular*}", "\\end{tabular}")
    s = re.sub(r"\\begin\{tabular\}\{(.*?)\}\n", lambda m: "\\begin{tabular}{" + simplify_colspec(m.group(1)) + "}\n", s)
    s = re.sub(r"\\includegraphics\[[^\]]*\]\{([^}]*)\}", r"\\includegraphics[width=16cm]{figures/\1.png}", s)
    for cmd in ("\\linenumbers", "\\FloatBarrier", "\\centering", "\\footnotesize", "\\scriptsize", "\\small",
                "\\raggedright", "\\allowbreak", "\\bottomrule", "\\toprule", "\\midrule"):
        s = s.replace(cmd + " ", " ").replace(cmd, "\\hline" if cmd in ("\\toprule", "\\midrule", "\\bottomrule") else "")
    s = re.sub(r"\\setlength\{[^}]*\}\{[^}]*\}", "", s)
    s = re.sub(r"\\cmidrule(\([a-z]*\))?\{[^}]*\}", "", s)
    s = re.sub(r"\\begin\{(table|figure)\*?\}\[[^\]]*\]", r"\\begin{\1}", s)
    s = re.sub(r"\\end\{(table|figure)\*\}", r"\\end{\1}", s)
    s = re.sub(r"%[^\n]*", lambda m: "" if not m.string[m.start() - 1:m.start()] == "\\" else m.group(0), s)

    # captions get their numbers; labels are dropped
    def number_caption(m):
        env, inner = m.group(1), m.group(2)
        lab = re.search(r"\\label\{([^}]*)\}", inner)
        num = labels.get(lab.group(1), "?") if lab else "?"
        i = inner.index("\\caption{")
        cap, end = braced(inner, i + 8)
        # Elsevier captions: "Table 1" on its own line above the table; "Fig. 1." run in below the figure
        head = f"\\textbf{{Table {num}}} TBLBRK " if env == "table" else f"\\textbf{{Fig. {num}.}} "
        inner = inner[:i] + "\\caption{" + head + cap + "}" + inner[end:]
        return f"\\begin{{{env}}}{inner}\\end{{{env}}}"
    s = re.sub(r"\\begin\{(table|figure)\}(.*?)\\end\{\1\}", number_caption, s, flags=re.S)

    # algorithms: the standalone rendering (figures/alg1_v1.png) with the PDF's number and caption
    def algorithm_block(m):
        inner = m.group(1)
        lab = re.search(r"\\label\{([^}]*)\}", inner)
        num = labels.get(lab.group(1), "?") if lab else "?"
        i = inner.index("\\caption{")
        cap, _ = braced(inner, i + 8)
        src = re.search(r"\\input\{figures/([^}]*)_body\}", inner).group(1)
        width = 16 if "*" in m.group(0)[:20] else 12
        return (f"\\begin{{figure}}\n\\includegraphics[width={width}cm]{{figures/{src}.png}}\n"
                f"\\caption{{\\textbf{{Algorithm {num}.}} {cap}}}\n\\end{{figure}}")
    s = re.sub(r"\\begin\{algorithm\*?\}(.*?)\\end\{algorithm\*?\}", algorithm_block, s, flags=re.S)

    # equations: native Word equations with the PDF's number
    def equation(m):
        inner = m.group(1)
        lab = re.search(r"\\label\{([^}]*)\}", inner)
        num = labels.get(lab.group(1), "?") if lab else ""
        inner = re.sub(r"\\label\{[^}]*\}", "", inner).strip().rstrip(",").rstrip(".")
        inner = re.sub(r"\\(begin|end)\{split\}", "", inner)              # one-line equations in Word
        inner = re.sub(r"\\\\\s*", " ", inner).replace("&", "")
        inner = inner.strip().rstrip(",").rstrip(".")
        for a, b in (("\\Bigl(", "\\left("), ("\\Bigr)", "\\right)"), ("\\bigl(", "\\left("), ("\\bigr)", "\\right)"),
                     ("\\textstyle", ""), ("b\\ \\mathrm{open}", "b \\text{ open}"), ("\\;", " "), ("\\,", " ")):
            inner = inner.replace(a, b)
        tail = m.group(2) or ""
        return f"\n\n\\[ {inner}{tail} \\qquad ({num}) \\]\n\n"
    s = re.sub(r"\\begin\{equation\}(.*?)([,.]?)\s*\\end\{equation\}", lambda m: equation(m), s, flags=re.S)

    # setting names such as $T300$, $c0$ as italic text, so they take the size of the text around them
    s = re.sub(r"\$([Tc])(\d+(?:\.\d+)?|\\infty)\$", lambda m: f"\\textit{{{m.group(1)}}}{m.group(2)}", s)

    # a lone minus sign as a text minus (a one-character equation widens table cells)
    s = s.replace("$-$", "\u2212")

    # superscript + and * as text, which every Word equation renderer accepts (LibreOffice drops "^+")
    s = re.sub(r"\^\{?\+\}?", r"^{\\text{+}}", s)
    s = re.sub(r"\^\{?\*\}?", r"^{\\text{*}}", s)

    # cross-references and citations
    s = re.sub(r"\\eqref\{([^}]*)\}", lambda m: f"({labels.get(m.group(1), '?')})", s)
    s = re.sub(r"\\ref\{([^}]*)\}", lambda m: labels.get(m.group(1), "?"), s)
    s = re.sub(r"\\citet\{([^}]*)\}", lambda m: cite_text(m.group(1), cites, textual=True), s)
    s = re.sub(r"~?\\citep?\{([^}]*)\}", lambda m: " " + cite_text(m.group(1), cites), s)
    s = re.sub(r"\\label\{[^}]*\}", "", s)

    # run-in paragraph headings, as in the PDF
    s = re.sub(r"\\paragraph\{([^}]*)\}\s*", lambda m: f"\\textit{{{m.group(1)}.}} ", s)

    # references from the bibliography of the LaTeX run
    bbl = BBL.read_text()
    bbl = re.sub(r"%[^\n]*\n", "", bbl)
    items = re.split(r"\\bibitem\[\{(?:[^{}]|\{[^{}]*\})*\}\]\{([^}]*)\}", bbl)
    refs = []
    for k in range(1, len(items), 2):
        key, text = items[k], items[k + 1]
        text = text.split("\\end{thebibliography}")[0].replace("\\newblock", " ")
        text = re.sub(r"\\bibinfo\{[^}]*\}", "", text)
        text = re.sub(r"\\href\s*\{([^}]*)\}\s*\{\\path\{([^}]*)\}\}", r"\\url{\1}", text)
        text = re.sub(r"\\DOIprefix\\doi\{([^}]*)\}", r"https://doi.org/\1", text)
        text = re.sub(r"\\doi\{([^}]*)\}", r"https://doi.org/\1", text)
        text = text.replace("\\URLprefix", "URL: ").replace("\\ArXivprefix", "arXiv:")
        text = re.sub(r"\s+", " ", text).strip()
        refs.append(f"REFMARK {text}")
    s = re.sub(r"\\bibliographystyle\{[^}]*\}\s*\\bibliography\{[^}]*\}",
               lambda m: "\\section*{References}\n\n" + "\n\n".join(refs) + "\n", s)
    return "\\documentclass{article}\n\\begin{document}\n" + s + "\n\\end{document}\n"


def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, val in kw.items():
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), val["val"])
        el.set(qn("w:sz"), str(val.get("sz", 4)))
        el.set(qn("w:color"), val.get("color", "000000"))


def fit_table(tbl, total_dxa=9072):
    """Column widths in proportion to their longest cell (merged cells excluded), on a fixed layout."""
    t = tbl._tbl
    rows = t.findall(qn("w:tr"))
    ncols = len(t.find(qn("w:tblGrid")).findall(qn("w:gridCol")))
    longest, word = [3] * ncols, [3] * ncols
    for tr in rows:
        col = 0
        for tc in tr.findall(qn("w:tc")):
            pr = tc.find(qn("w:tcPr"))
            span = pr.find(qn("w:gridSpan")) if pr is not None else None
            n = int(span.get(qn("w:val"))) if span is not None else 1
            text = "".join(x.text or "" for x in tc.iter(qn("w:t")))
            if n == 1 and col < ncols:
                longest[col] = max(longest[col], min(len(text) + (2 if "[" in text else 0), 40))
                word[col] = max([word[col]] + [min(len(w), 16) for w in text.split()])
            col += n
    weights = [max(l, w) + 2 for l, w in zip(longest, word)]
    widths = [int(total_dxa * w / sum(weights)) for w in weights]
    grid = t.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    tblPr = t.find(qn("w:tblPr"))
    for tag in ("w:tblW", "w:tblLayout"):
        el = tblPr.find(qn(tag))
        if el is not None:
            tblPr.remove(el)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(sum(widths))); tw.set(qn("w:type"), "dxa"); tblPr.append(tw)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    mar = tblPr.find(qn("w:tblCellMar"))
    if mar is not None:
        tblPr.remove(mar)
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), "57"); el.set(qn("w:type"), "dxa"); mar.append(el)
    tblPr.append(mar)
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
             "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
    kids = list(tblPr)
    for k in kids:
        tblPr.remove(k)
    for k in sorted(kids, key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 99):
        tblPr.append(k)
    for tr in rows:
        col = 0
        for tc in tr.findall(qn("w:tc")):
            pr = tc.find(qn("w:tcPr"))
            if pr is None:
                pr = OxmlElement("w:tcPr"); tc.insert(0, pr)
            span = pr.find(qn("w:gridSpan"))
            n = int(span.get(qn("w:val"))) if span is not None else 1
            w = sum(widths[col:col + n])
            el = pr.find(qn("w:tcW"))
            if el is None:
                el = OxmlElement("w:tcW"); pr.insert(0, el)
            el.set(qn("w:w"), str(w)); el.set(qn("w:type"), "dxa")
            col += n


def style_document(path):
    doc = Document(path)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    pgmar = sec._sectPr.find(qn("w:pgMar"))
    for a, v in (("w:header", "708"), ("w:footer", "708"), ("w:gutter", "0")):
        if pgmar.get(qn(a)) is None:
            pgmar.set(qn(a), v)
    sizes = {"Normal": 12, "BodyText": 12, "FirstParagraph": 12, "Compact": 10, "Heading1": 14, "Heading2": 12,
             "Heading3": 12, "Heading4": 12, "TableCaption": 10, "ImageCaption": 10, "CaptionedFigure": 12,
             "Title": 16, "Abstract": 11}
    for st in doc.styles.element.findall(qn("w:style")):
        sid = st.get(qn("w:styleId"))
        if st.get(qn("w:type")) not in ("paragraph", "character"):
            continue
        rpr = st.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            after = [c for c in st if c.tag.split("}")[1] in ("tblPr", "trPr", "tcPr", "tblStylePr")]
            if after:
                after[0].addprevious(rpr)
            else:
                st.append(rpr)
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
        for a in list(rf.attrib):
            if a.endswith("Theme") or a.endswith("theme"):
                del rf.attrib[a]
        if sid not in ("SourceCode", "VerbatimChar"):
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rf.set(qn(a), FONT)
        col = rpr.find(qn("w:color"))
        if col is not None:
            rpr.remove(col)
        if sid in sizes:
            for tag in ("w:sz", "w:szCs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = OxmlElement(tag); rpr.append(el)
                el.set(qn("w:val"), str(2 * sizes[sid]))
        if sid and sid.startswith("Heading"):
            if rpr.find(qn("w:b")) is None:
                rpr.append(OxmlElement("w:b"))
            if sid == "Heading2" and rpr.find(qn("w:i")) is None:
                rpr.append(OxmlElement("w:i"))
            ppr = st.find(qn("w:pPr"))
            sp = ppr.find(qn("w:spacing")) if ppr is not None else None
            if sp is not None:
                sp.set(qn("w:before"), "240"); sp.set(qn("w:after"), "120")
    for name in ("Body Text", "First Paragraph", "Normal"):
        for st in doc.styles:
            if st.name == name:
                st.paragraph_format.line_spacing = 1.5
                st.paragraph_format.space_after = Pt(6)
    for name in ("Table Caption", "Image Caption"):
        for st in doc.styles:
            if st.name == name:
                st.font.italic = False
                st.paragraph_format.space_after = Pt(4)
                st.paragraph_format.line_spacing = 1.0

    marks = {"TITLEMARK ": ("title", 16), "AUTHORMARK ": ("author", 12), "AFFILMARK ": ("affil", 11),
             "CORRMARK ": ("corr", 10), "ABSTRACTHEAD ": ("abshead", 12), "ABSTRACTMARK ": ("abstract", 12),
             "KEYWORDSMARK ": ("keywords", 12), "REFMARK ": ("ref", 10)}
    for p in doc.paragraphs:
        for mark, (kind, size) in marks.items():
            if p.text.startswith(mark):
                for r in p.runs:                                       # remove the marker text
                    if mark.strip() in r.text:
                        r.text = r.text.replace(mark, "").replace(mark.strip(), "")
                        break
                for r in p.runs:                                       # and any space it left behind
                    if r.text:
                        r.text = r.text.lstrip()
                        break
                for r in p._p.iter(qn("w:r")):                           # runs inside hyperlinks too
                    rpr = r.find(qn("w:rPr"))
                    if rpr is None:
                        rpr = OxmlElement("w:rPr"); r.insert(0, rpr)
                    for tag in ("w:sz", "w:szCs"):
                        el = rpr.find(qn(tag))
                        if el is None:
                            el = OxmlElement(tag); rpr.append(el)
                        el.set(qn("w:val"), str(2 * size))
                for r in p.runs:
                    if kind in ("title", "abshead"):
                        r.bold = True
                    if kind == "affil":
                        r.italic = True
                pf = p.paragraph_format
                if kind in ("title", "author", "affil", "corr"):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if kind != "corr" else WD_ALIGN_PARAGRAPH.LEFT
                    pf.space_after = Pt(6)
                if kind == "ref":
                    pf.left_indent, pf.first_line_indent = Cm(0.8), Cm(-0.8)
                    pf.line_spacing, pf.space_after = 1.0, Pt(3)
                if kind == "abstract":
                    pf.line_spacing = 1.5
                break
    for p in doc.paragraphs:                       # "Table n" on its own line above its caption text
        for r in p.runs:
            if "TBLBRK" in r.text:
                before, after = r.text.split("TBLBRK", 1)
                r.text = before.rstrip()
                br = OxmlElement("w:br")
                r._r.append(br)
                t = OxmlElement("w:t"); t.text = after.lstrip(); t.set(qn("xml:space"), "preserve")
                r._r.append(t)
                break
    for tbl in doc.tables:
        fit_table(tbl)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        n = len(tbl.rows)
        trs = tbl._tbl.findall(qn("w:tr"))
        first_cells = ["".join(x.text or "" for x in tr.find(qn("w:tc")).iter(qn("w:t"))).strip() for tr in trs]
        texts = ["".join(x.text or "" for x in tr.iter(qn("w:t"))) for tr in trs]
        marked = [i for i, tr in enumerate(trs) if tr.find(qn("w:trPr")) is not None and tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None]
        if marked:
            head_end = max(marked)
        else:                                  # multi-row headers: up to the row naming the first column
            head_end, seen = 0, False
            for i, fc in enumerate(first_cells):
                if not seen and fc:
                    seen, head_end = True, i
                elif seen and not fc and not re.search(r"\d", texts[i]):
                    head_end = i
                elif seen:
                    break
        for i, row in enumerate(tbl.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                edges = {}
                if i == 0:
                    edges["top"] = {"val": "single", "sz": 8}
                if i == head_end and i != n - 1:
                    edges["bottom"] = {"val": "single", "sz": 4}
                if i == n - 1:
                    edges["bottom"] = {"val": "single", "sz": 8}
                if edges:
                    set_cell_border(cell, **edges)
    doc.save(path)


def main():
    if not AUX.exists() or not BBL.exists():
        sys.exit("run latexmk -pdf main.tex first (main.aux and main.bbl are needed)")
    tex = body_latex()
    with tempfile.TemporaryDirectory() as d:
        src = pathlib.Path(d) / "word.tex"
        src.write_text(tex)
        (HERE / "build_word.tex").write_text(tex)                    # kept for inspection; not committed
        subprocess.run(["pandoc", str(src), "-f", "latex", "-t", "docx", "--number-sections",
                        "--resource-path", str(HERE), "-o", str(OUT)], check=True)
    style_document(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
